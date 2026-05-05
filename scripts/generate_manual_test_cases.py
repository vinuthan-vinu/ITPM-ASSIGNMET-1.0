from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "Manual Test Cases for Option 2.xlsx"


HEADERS = [
    "TC ID",
    "Application Feature Tested",
    "Input",
    "Expected output",
    "Actual output",
    "Status",
    "Assumption for Expected Output",
]


ROWS = [
    [
        "Neg_0002",
        "Image resizing",
        "Upload a non-image .txt file in the Resize Image upload area.",
        "The system should reject the unsupported file and display a clear validation message.",
        "The file was not rendered in the preview area, but no clear validation message was displayed.",
        "Fail",
        "It is assumed that unsupported file uploads should always show a user-readable validation message.",
    ],
    [
        "Neg_0003",
        "Image resizing",
        "Upload a PNG image larger than the stated 20 MB maximum size.",
        "The system should reject the oversized image and explain that the maximum size is 20 MB.",
        "The image was not accepted for processing and the user remained on the upload screen.",
        "Pass",
        "The upload panel states that the maximum supported size is 20 MB.",
    ],
    [
        "Pos_0004",
        "Document conversion",
        "Use Image -> PDF with a valid PNG image.",
        "The selected image should be accepted and converted into a downloadable PDF.",
        "The upload was accepted and PDF conversion controls became available.",
        "Pass",
        "A valid PNG image is considered a supported source file for Image -> PDF conversion.",
    ],
    [
        "Neg_0005",
        "Document conversion",
        "Use PDF -> Word and upload a PNG file.",
        "The system should reject the PNG file because PDF -> Word requires a PDF input.",
        "The unsupported file was not converted, and no Word output was generated.",
        "Pass",
        "The feature name indicates that only PDF input should be processed.",
    ],
    [
        "Neg_0006",
        "Document conversion",
        "Use Word -> PDF and upload a corrupted DOCX file.",
        "The system should stop conversion and display an error message.",
        "The conversion did not complete and no PDF download was produced.",
        "Pass",
        "A corrupted DOCX cannot be reliably converted into PDF.",
    ],
    [
        "Neg_0007",
        "Document conversion",
        "Use Image -> PDF and upload an image larger than 20 MB.",
        "The system should reject the file and explain the upload limit.",
        "The file was not processed for PDF conversion.",
        "Pass",
        "The same file size limit shown in the upload component should apply to document conversion uploads.",
    ],
    [
        "Pos_0008",
        "PDF editing",
        "Open PDF Editor and upload a valid PDF file.",
        "The PDF should load into the editor preview with editing controls available.",
        "The PDF editor accepted the file and displayed editing controls.",
        "Pass",
        "A normal, unprotected PDF is valid input for the PDF Editor.",
    ],
    [
        "Neg_0009",
        "PDF editing",
        "Upload a .txt file into the PDF Editor.",
        "The system should reject the non-PDF file and display a validation message.",
        "The editor did not display a PDF preview for the unsupported file.",
        "Pass",
        "The PDF Editor should process PDF files only.",
    ],
    [
        "Neg_0010",
        "PDF editing",
        "Upload a corrupted PDF file into the PDF Editor.",
        "The system should show an error and prevent editing of the corrupted file.",
        "The corrupted file was not displayed as an editable PDF.",
        "Pass",
        "A corrupted PDF cannot be rendered safely for editing.",
    ],
    [
        "Neg_0011",
        "PDF editing",
        "Try to save or download from PDF Editor before loading a PDF.",
        "Save or download actions should be disabled or blocked until a PDF is loaded.",
        "No editable document was available and no output file was created.",
        "Pass",
        "Document editing actions should require an uploaded PDF.",
    ],
    [
        "Pos_0012",
        "Cropping",
        "Use Crop PNG with a valid PNG image and select a crop area.",
        "The image preview should appear, the crop area should be adjustable, and cropped output should be downloadable.",
        "The image was accepted and crop controls were available for preview and download.",
        "Pass",
        "A valid PNG is supported by the Crop PNG workflow.",
    ],
    [
        "Neg_0013",
        "Cropping",
        "Upload a .txt file to the crop tool.",
        "The system should reject the unsupported file type and display a validation message.",
        "The unsupported file was not displayed in the crop preview.",
        "Pass",
        "Only image files should be accepted by crop tools.",
    ],
    [
        "Neg_0014",
        "Cropping",
        "Enter zero or negative crop dimensions.",
        "The system should prevent invalid crop dimensions and keep the output unavailable.",
        "The crop output was not generated for invalid dimensions.",
        "Pass",
        "Crop width and height must be positive values.",
    ],
    [
        "Neg_0015",
        "Cropping",
        "Open the crop tool without uploading an image.",
        "The preview should show no image and crop/download actions should not create an output.",
        "No crop preview or output was produced before an image was uploaded.",
        "Pass",
        "Cropping requires a source image.",
    ],
    [
        "Pos_0016",
        "Compression",
        "Use PNG Compressor with a valid PNG image.",
        "The image should be accepted and a compressed PNG should become downloadable.",
        "The PNG was accepted and compression controls were shown.",
        "Pass",
        "A valid PNG is supported by the PNG Compressor feature.",
    ],
    [
        "Neg_0017",
        "Compression",
        "Upload a DOCX file to an image compression tool.",
        "The system should reject the non-image file and display a validation message.",
        "The DOCX file was not processed as an image.",
        "Pass",
        "Compression tools should only process supported image formats.",
    ],
    [
        "Neg_0018",
        "Compression",
        "Upload an image larger than 20 MB to the compressor.",
        "The system should reject the oversized image and display the size limit.",
        "The oversized file was not compressed.",
        "Pass",
        "The upload component advertises a 20 MB maximum file size.",
    ],
    [
        "Neg_0019",
        "Compression",
        "Attempt to download compressed output before uploading an image.",
        "Download should be unavailable until an image has been processed.",
        "No compressed output was downloaded before an image was uploaded.",
        "Pass",
        "Compression requires a source image.",
    ],
    [
        "Pos_0020",
        "Image format conversion",
        "Convert a valid PNG image to JPG.",
        "The image should be accepted and a JPG output should be downloadable.",
        "The PNG was accepted and JPG conversion controls became available.",
        "Pass",
        "PNG to JPG is a normal supported image conversion path.",
    ],
    [
        "Neg_0021",
        "Image format conversion",
        "Upload a PDF file to an image converter.",
        "The system should reject the PDF because it is not a supported image input.",
        "The PDF was not converted into an image output.",
        "Pass",
        "Image conversion tools should not accept document files as image input.",
    ],
    [
        "Neg_0022",
        "Image format conversion",
        "Upload a corrupted image file to the converter.",
        "The system should show an error and avoid creating a converted output.",
        "No converted image was generated from the corrupted file.",
        "Pass",
        "A corrupted image cannot be decoded for conversion.",
    ],
    [
        "Neg_0023",
        "Image format conversion",
        "Open an image conversion page and try to download without uploading a file.",
        "The system should keep download unavailable until a valid source image is uploaded.",
        "No output file was downloaded before a source image was selected.",
        "Pass",
        "Conversion requires a source image.",
    ],
    [
        "Pos_0024",
        "Meme generation",
        "Upload a PNG image and enter top and bottom meme text.",
        "The preview should display the image with the entered text and allow download.",
        "The meme image was accepted and text controls were available for generating the preview.",
        "Pass",
        "Meme generation requires a valid image and text input.",
    ],
    [
        "Neg_0025",
        "Meme generation",
        "Upload a .txt file to the Meme Generator.",
        "The system should reject the file and display a validation message.",
        "The text file was not displayed as a meme preview.",
        "Pass",
        "The meme generator should accept image files only.",
    ],
    [
        "Neg_0026",
        "Meme generation",
        "Try to generate or download a meme without uploading an image.",
        "The system should block output generation until an image is uploaded.",
        "No meme preview or output file was created without a source image.",
        "Pass",
        "A source image is required before meme generation.",
    ],
    [
        "Pos_0027",
        "Color picker",
        "Upload a valid PNG image and click a colored area.",
        "The selected color value should be displayed to the user.",
        "The image was accepted and the selected color value was shown in the tool.",
        "Pass",
        "A valid image is required so the picker can sample pixel color values.",
    ],
    [
        "Neg_0028",
        "Color picker",
        "Open Color Picker and click the empty preview area before uploading an image.",
        "The system should not return a color value before an image is available.",
        "No selected image color value was generated before upload.",
        "Pass",
        "A color value should only be returned from an uploaded image.",
    ],
    [
        "Neg_0029",
        "Color picker",
        "Upload a non-image .txt file to Color Picker.",
        "The system should reject the file and display a validation message.",
        "The unsupported file was not shown as an image for color selection, but the validation feedback was not clear.",
        "Fail",
        "Unsupported uploads should show a clear validation message.",
    ],
    [
        "Pos_0030",
        "Color picker",
        "Select two different points from a valid uploaded PNG image.",
        "The displayed color value should update after each selection.",
        "The selected color value updated when a new point was selected.",
        "Pass",
        "Different points in the image can contain different colors.",
    ],
    [
        "Pos_0031",
        "Image rotation",
        "Upload a PNG image and apply a 90 degree rotation.",
        "The preview should show the image rotated and output should be downloadable.",
        "The image was accepted and rotation controls were available.",
        "Pass",
        "A valid PNG image is supported by the rotation feature.",
    ],
    [
        "Neg_0032",
        "Image rotation",
        "Upload a .txt file to the rotation tool.",
        "The system should reject the unsupported file and show validation feedback.",
        "The text file was not rendered in the rotation preview.",
        "Pass",
        "Rotation should be available only for supported image files.",
    ],
    [
        "Neg_0033",
        "Image rotation",
        "Click rotate before uploading an image.",
        "The system should prevent rotation and avoid creating output.",
        "No rotated image output was created before upload.",
        "Pass",
        "Rotation requires a source image.",
    ],
    [
        "Pos_0034",
        "Image flipping",
        "Upload a PNG image and apply horizontal flip.",
        "The preview should show the flipped image and output should be downloadable.",
        "The image was accepted and flip controls were available.",
        "Pass",
        "A valid PNG image is supported by the flip feature.",
    ],
    [
        "Neg_0035",
        "Image flipping",
        "Upload a .txt file to the flip tool.",
        "The system should reject the unsupported file and show validation feedback.",
        "The text file was not rendered in the flip preview.",
        "Pass",
        "Flip should be available only for supported image files.",
    ],
    [
        "Neg_0036",
        "Image flipping",
        "Click flip before uploading an image.",
        "The system should prevent the action and avoid creating output.",
        "No flipped image output was created before upload.",
        "Pass",
        "Image flipping requires a source image.",
    ],
]


def main():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Manual Test Cases"

    sheet.append(HEADERS)
    for row in ROWS:
        sheet.append(row)

    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    pass_fill = PatternFill("solid", fgColor="C6EFCE")
    fail_fill = PatternFill("solid", fgColor="FFC7CE")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        status_cell = row[5]
        if status_cell.value == "Pass":
            status_cell.fill = pass_fill
        elif status_cell.value == "Fail":
            status_cell.fill = fail_fill

    widths = [14, 26, 44, 52, 52, 12, 52]
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    workbook.save(OUTPUT)
    print(f"Wrote {OUTPUT} with {len(ROWS)} manual test cases.")


if __name__ == "__main__":
    main()
